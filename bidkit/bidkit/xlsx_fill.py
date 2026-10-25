#!/usr/bin/env python3
"""
xlsx_fill.py — write Vendor Response + Comments (or any cell values) into a COPY of an agency
workbook by SURGICAL XML EDIT, preserving everything Excel needs.

Why (three traps, all hit on a prior bid):
  1. openpyxl.save() strips the x14 dropdown data-validations agencies use for the controlled
     response vocabulary. So we never save through openpyxl — we edit the sheet XML inside the
     zip and re-zip.
  2. Excel requires a row's <c> cells in ascending column order. Appending C/D to a row that
     already has later columns makes Excel "repair" the file and STRIP the row's data on open
     (blanked whole tabs). Every write re-sorts the row's cells.
  3. In a merged range only the top-left anchor holds a value. Writing into a merged-away cell
     hides the value and can trigger the repair warning. Every target is checked against the
     sheet's <mergeCells>; merged-away targets are REFUSED and reported (never silently dropped).

Two modes:
  responses mode  — python3 -m bidkit.xlsx_fill <opportunity.json> <wb-key> <responses.json>
                    responses.json = list of {tab,row,response,comment} (xlsx_extract output,
                    filled in). Response values are validated against the workbook's control-tab
                    vocabulary (config workbooks.<key>.control_tab/control_range).
  cells mode      — python3 -m bidkit.xlsx_fill <opportunity.json> <wb-key> --cells <cells.json>
                    cells.json = {"<Sheet>": {"B4": "text", ...}} for form-style workbooks.

Output: <draft>/<original filename>  (a copy; the rfp/ original is never modified)
Exit 1 if any write was refused or any response value is outside the vocabulary.
"""
import json, os, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl
from . import config
from ._xlsx_common import (read_package, write_package, sheet_xml_map, merge_map, set_cell,
                           num_to_col, count_validations, out_of_order_rows)


def load_vocabulary(path, control_tab, control_range):
    """Read the allowed dropdown strings from the control tab. Returns a set (empty = no check)."""
    if not control_tab or not control_range:
        return set()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if control_tab not in wb.sheetnames:
        wb.close(); return set()
    ws = wb[control_tab]
    vocab = set()
    for row in ws[control_range]:
        for c in row:
            if c.value is not None and str(c.value).strip():
                vocab.add(str(c.value).strip())
    wb.close()
    return vocab


class FillResult:
    def __init__(self):
        self.written = 0
        self.refused_merged = []     # (sheet, ref, reason) — a VALUE had nowhere valid to go
        self.skipped_comments = []   # (sheet, ref) — comment cell merged into the response cell; informational
        self.bad_vocab = []          # (sheet, row, value)
        self.missing_sheets = []
        self.validations_before = 0
        self.validations_after = 0
        self.out_of_order_after = 0

    @property
    def ok(self):
        return not (self.refused_merged or self.bad_vocab or self.missing_sheets
                    or self.out_of_order_after
                    or self.validations_after != self.validations_before)

    def report(self):
        lines = [f"written: {self.written}",
                 f"x14 validations before/after: {self.validations_before}/{self.validations_after}",
                 f"out-of-order rows after: {self.out_of_order_after}"]
        for s, ref, why in self.refused_merged:
            lines.append(f"REFUSED merged-away write {s}!{ref}: {why}")
        for s, ref in self.skipped_comments:
            lines.append(f"note: comment skipped at {s}!{ref} (merged into the response cell; no separate comment cell)")
        for s, r, v in self.bad_vocab:
            lines.append(f"REFUSED bad vocabulary {s} row {r}: {v!r}")
        for s in self.missing_sheets:
            lines.append(f"MISSING sheet {s!r}")
        return "\n".join(lines)


def fill_cells(src, dst, sheet_cells, vocab=None, vocab_cols=None, soft_refs=None):
    """sheet_cells: {sheet: {ref: text}}. vocab/vocab_cols: if given, values written into any
    column in vocab_cols must be in vocab. soft_refs: {sheet: {ref,...}} of secondary cells
    (comments) that are SKIPPED (informational) rather than REFUSED when merged-away — because
    on form-style rows the comment column is merged into the response cell and the value has
    a legitimate home in the anchor. Primary values in merged-away cells are always refused.
    Returns FillResult."""
    res = FillResult()
    soft_refs = soft_refs or {}
    names, contents = read_package(src)
    smap = sheet_xml_map(contents)
    res.validations_before = sum(count_validations(contents[p].decode("utf-8"))[0]
                                 for p in smap.values() if p in contents)
    for sheet, cells in sheet_cells.items():
        sp = smap.get(sheet)
        if not sp or sp not in contents:
            res.missing_sheets.append(sheet); continue
        xml = contents[sp].decode("utf-8")
        anchors, away = merge_map(xml)
        for ref, text in cells.items():
            from ._xlsx_common import split_ref
            col, row = split_ref(ref)
            if (col, row) in away:
                if ref in soft_refs.get(sheet, set()):
                    res.skipped_comments.append((sheet, ref))
                else:
                    res.refused_merged.append((sheet, ref, "cell is inside a merged range but is not its anchor"))
                continue
            if vocab and vocab_cols and col in vocab_cols and text not in vocab:
                res.bad_vocab.append((sheet, row, text)); continue
            xml = set_cell(xml, ref, text); res.written += 1
        contents[sp] = xml.encode("utf-8")
    # post-checks before writing
    res.validations_after = sum(count_validations(contents[p].decode("utf-8"))[0]
                                for p in smap.values() if p in contents)
    res.out_of_order_after = sum(len(out_of_order_rows(contents[p].decode("utf-8")))
                                 for p in smap.values() if p in contents)
    os.makedirs(os.path.dirname(dst) or ".", exist_ok=True)
    write_package(dst, names, contents)
    return res


def fill_responses(opp, key, responses):
    wbcfg = opp.workbook(key)
    src = opp.resolve(wbcfg["path"])
    dst = os.path.join(opp.folder("draft", create=True), os.path.basename(src))
    cols = wbcfg["cols"]
    resp_col = num_to_col(int(cols["resp"])); com_col = num_to_col(int(cols["comment"]))
    vocab = load_vocabulary(src, wbcfg.get("control_tab"), wbcfg.get("control_range"))
    sheet_cells, soft = {}, {}
    for item in responses:
        resp = (item.get("response") or "").strip()
        if not resp:
            continue
        row = int(item["row"])
        d = sheet_cells.setdefault(item["tab"], {})
        d[f"{resp_col}{row}"] = resp
        com = (item.get("comment") or "").strip()
        if com:
            d[f"{com_col}{row}"] = com
            soft.setdefault(item["tab"], set()).add(f"{com_col}{row}")
    res = fill_cells(src, dst, sheet_cells, vocab=vocab, vocab_cols={resp_col}, soft_refs=soft)
    return res, dst


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cfg", help="opportunity.json"); ap.add_argument("wb", help="workbook key, e.g. A1")
    ap.add_argument("responses", nargs="?", help="responses.json (xlsx_extract output with response/comment filled)")
    ap.add_argument("--cells", help='cells.json {"<Sheet>": {"B4": "text", ...}} for form-style workbooks')
    a = ap.parse_args(argv)
    if not a.responses and not a.cells:
        ap.error("give responses.json or --cells cells.json")
    opp = config.load(a.cfg); key = a.wb
    if a.cells:
        with open(a.cells, encoding="utf-8") as fh:
            sheet_cells = json.load(fh)
        wbcfg = opp.workbook(key); src = opp.resolve(wbcfg["path"])
        dst = os.path.join(opp.folder("draft", create=True), os.path.basename(src))
        res = fill_cells(src, dst, sheet_cells)
    else:
        with open(a.responses, encoding="utf-8") as fh:
            responses = json.load(fh)
        res, dst = fill_responses(opp, key, responses)
    print(f"{key}: -> {dst}")
    print(res.report())
    return 0 if res.ok else 1


if __name__ == "__main__":
    sys.exit(main())
