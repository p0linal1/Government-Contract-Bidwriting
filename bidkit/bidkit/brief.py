#!/usr/bin/env python3
"""
brief.py — turn a filed opportunity (email + rfp/ documents) into a FACTS SHEET and a RESEARCH PLAN.

Why: intake starts the same way for every solicitation regardless of response format — find the
dates, the money signals, the contacts, the format rules, the required attachments, the evaluation
criteria — and then research the incumbent and prior contract to sanity-check the value. On a prior
bid, finding the incumbent's previous contract (a council packet PDF found by web search) was the
single most valuable input to pricing. This module does the deterministic half:

  opportunity-facts.md / .json   every extracted fact WITH its source snippet, labelled DETECTED
                                 (nothing here is confirmed until a human reads the RFP)
  research/research-plan.md      the exact web-search queries to run for this agency/system, and
                                 what to look for (incumbent, prior award $, term, council packets,
                                 competitor awards); research/findings.md template to fill
  opportunity.json               due/questions/prebid dates set if empty (confirmed:false),
                                 format_detected block (never overwrites the authoritative format block)

The narrative summary and the estimated value are judgment calls made in the Kiro intake step
("run intake on <opp>") with this sheet + the research findings in hand. This tool never guesses
them; it surfaces every signal so the human/agent doesn't have to hunt.

Inputs : <opportunity-dir> with initial-email/email.md and/or rfp/*.pdf|xlsx|docx|txt|md|html
CLI    : python3 -m bidkit.brief <opportunity-dir> [--json] [--no-research] [--max-snippets N]
"""
import argparse, datetime as dt, json, os, re, sys, warnings
warnings.filterwarnings("ignore")
from .inbox import DATE_PATTERNS, MONTHS, guess_due_date  # shared date grammar

DOC_EXT = (".pdf", ".xlsx", ".xlsm", ".xls", ".docx", ".txt", ".md", ".html", ".htm")

# cue word -> event kind (first match wins, checked in order)
EVENT_CUES = [
    ("questions", re.compile(r"\b(questions?|inquir(?:y|ies)|clarification)\b", re.I)),
    ("prebid", re.compile(r"\b(pre-?(?:bid|proposal|submittal)|site visit|walk-?through|(?:vendor|bidders?) conference)\b", re.I)),
    ("addenda", re.compile(r"\b(addend(?:um|a)|amendment)\b", re.I)),
    ("award", re.compile(r"\b(award|intent to award|notice of intent|selection|shortlist|finalist)\b", re.I)),
    ("interviews", re.compile(r"\b(interview|presentation|demonstration|demo)\b", re.I)),
    ("issued", re.compile(r"\b(issued|released?|posted|date of (?:this )?(?:rfp|rfq|solicitation)|advertis)\b", re.I)),
    ("start", re.compile(r"\b(contract (?:start|begin|commence)|commencement|go-?live|notice to proceed|effective date)\b", re.I)),
    ("due", re.compile(r"\b(due|deadline|clos(?:e|es|ing)|submit(?:ted|tal|ssion)s?|must be received|no later than|receipt of proposals)\b", re.I)),
]
TIME_RE = re.compile(r"\b(\d{1,2})(?::(\d{2}))?\s*([ap])\.?m\.?\b", re.I)
TZ_RE = re.compile(r"\b(P[SD]?T|Pacific(?: Time| Standard Time| Daylight Time)?|M[SD]?T|Mountain|C[SD]?T|Central|E[SD]?T|Eastern|local time)\b", re.I)
MONEY_RE = re.compile(r"\$\s?(\d{1,3}(?:,\d{3})+|\d+)(?:\.(\d{2}))?\s*(million|mm|m|k|thousand|billion)?\b", re.I)
MONEY_CUES = [
    ("value", re.compile(r"\b(budget\w*|estimat\w*|anticipat\w*|not[- ]to[- ]exceed|nte|maximum|ceiling|fund(?:s|ing|ed)?|appropriat\w*|contract (?:value|amount|price)|total (?:cost|value|price)|award\w*|allocat\w*)\b", re.I)),
    ("bond", re.compile(r"\b(bond|surety)\b", re.I)),
    ("insurance", re.compile(r"\b(insurance|liabilit|coverage|e&o|errors and omissions|cyber)\b", re.I)),
    ("penalty", re.compile(r"\b(liquidated|penalt|damages|fee)\b", re.I)),
]
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\(\d{3}\)\s?|\b\d{3}[.\-\s])\d{3}[.\-\s]\d{4}\b")
CONTACT_LINE_RE = re.compile(r"\b(coordinator|contact|buyer|procurement|purchasing|contracting officer|point of contact|inquiries)\b", re.I)
FORMAT_RE = re.compile(r"\b(font|\d{1,2}[- ]?(?:point|pt)\b|margins?|page limit|maximum of .{0,20}pages|pages? (?:in )?(?:length|total)|8\s?[½1/2]|8\.5|single[- ]spaced|double[- ]spaced|arial|times new roman|calibri|tabs?|numbered pages|file size|mb\b)", re.I)
_LABEL_WORDS = r"(?:Attachment|Exhibit|Appendix|Schedule|Form|Addendum|Enclosure)"
# label + optional "- Title": the title stops at newline / . ; , and cannot run into the NEXT label word
ATTACH_RE = re.compile(r"\b(" + _LABEL_WORDS + r")\s+((?:No\.?\s*)?[A-Z]{1,2}-?\d{0,2}|\d{1,2}|[IVX]{1,4})\b"
                       r"(?:\s*[-–—:]\s*((?:(?!\b" + _LABEL_WORDS + r"\b)[^\n.;,]){3,70}))?")
NUMWORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
            "fifteen": 15, "twenty": 20, "twenty-four": 24, "twenty-five": 25, "thirty": 30, "forty": 40, "fifty": 50}
EVAL_RE = re.compile(r"(\d{1,3})\s*(points?|pts|%|percent)\b", re.I)
EVAL_CUE_RE = re.compile(r"\b(evaluat|criteria|scor|weight|rated|rank)\b", re.I)
SUBMIT_RE = re.compile(r"\b(submi\w*|deliver\w*|received?)\b", re.I)   # submit/submitted/submittals/submission …
SUBMIT_HOW_RE = re.compile(r"\b(e-?mail|electronic(?:ally)?|portal|upload|hard cop(?:y|ies)|sealed|usb|flash drive|original|copies|pdf|subject line)\b", re.I)
SETASIDE_RE = re.compile(r"\b(small business|minority|women[- ]owned|veteran|omwbe|mwbe|dbe|wbe|mbe|local preference|section 3|hubzone|8\(a\)|disadvantaged)\b", re.I)
TERM_RE = re.compile(r"\b(?:(one|two|three|four|five|six|seven|eight|nine|ten|\d{1,2})\s*(?:\(\d+\)\s*)?[- ]?(year|month)s?)\b.{0,40}?\b(term|duration|period|renew|option|extension)\b|\b(term|duration|period|renew|option|extension)\b.{0,40}?\b(one|two|three|four|five|six|seven|eight|nine|ten|\d{1,2})\s*(?:\(\d+\)\s*)?[- ]?(year|month)s?\b", re.I)
TITLE_STOP = set("the a an and of for to in on by with request proposal proposals rfp rfq rfi solicitation".split())


# ----------------------------------------------------------------------------- text extraction
def extract_text(path):
    """Text from pdf/xlsx/docx/txt/md/html. Returns (text, meta) where meta has pages/sheets."""
    ext = os.path.splitext(path)[1].lower()
    meta = {"file": os.path.basename(path), "type": ext.lstrip(".")}
    try:
        if ext == ".pdf":
            import fitz
            d = fitz.open(path); parts = []
            for i in range(d.page_count):
                parts.append(d[i].get_text())
            meta["pages"] = d.page_count; d.close()
            text = "\n".join(parts)
        elif ext in (".xlsx", ".xlsm", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True); parts = []
            for ws in wb.worksheets:
                parts.append(f"## sheet {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None and str(c).strip()]
                    if cells: parts.append(" | ".join(cells))
            meta["sheets"] = len(wb.worksheets); wb.close()
            text = "\n".join(parts)
        elif ext == ".docx":
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            text = "\n".join(parts)
        elif ext in (".html", ".htm"):
            raw = open(path, encoding="utf-8", errors="ignore").read()
            raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
            text = re.sub(r"<[^>]+>", " ", raw)
            import html as _h; text = _h.unescape(text)
        else:
            text = open(path, encoding="utf-8", errors="ignore").read()
    except Exception as e:
        meta["error"] = str(e)[:160]; text = ""
    text = text.replace("\xa0", " ").replace("\u00ad", "")
    meta["chars"] = len(text); meta["words"] = len(text.split())
    return text, meta


def gather_sources(opp_dir):
    """[(label, path)] for every readable document: email.md first, then rfp/ (recursively)."""
    out = []
    em = os.path.join(opp_dir, "initial-email", "email.md")
    if os.path.exists(em): out.append(("email", em))
    fu = os.path.join(opp_dir, "initial-email", "followups")
    if os.path.isdir(fu):
        for d in sorted(os.listdir(fu)):
            p = os.path.join(fu, d, "email.md")
            if os.path.exists(p): out.append((f"followup:{d}", p))
    rfp = os.path.join(opp_dir, "rfp")
    if os.path.isdir(rfp):
        for root, _, files in os.walk(rfp):
            for f in sorted(files):
                # portal-*.{html,md,json} are bidkit.portal's own snapshots/reports, not solicitation text
                if f.lower().endswith(DOC_EXT) and not f.startswith("portal-"):
                    out.append((os.path.relpath(os.path.join(root, f), rfp), os.path.join(root, f)))
    return out


def load_portal_records(opp_dir):
    """rfp/portal-record.json written by bidkit.portal: what the portal's own solicitation page states
    (reference number, organization, posted/closed dates, estimated value, contact, file list). [] if absent."""
    p = os.path.join(opp_dir, "rfp", "portal-record.json")
    if not os.path.exists(p):
        return []
    try:
        recs = json.load(open(p, encoding="utf-8"))
        return [r for r in recs if isinstance(r, dict)] if isinstance(recs, list) else []
    except Exception:
        return []


INVERTED_ORG = re.compile(r"^(?P<name>.+?),\s*(?P<head>(?:City|Town|County|Port|Village|State|University|College|Office|Dept\.?|Department|Board|Bureau|Division)\s+of(?:\s+the)?)$", re.I)


def normalize_org_name(name):
    """Portals list organizations index-style: 'Springfield, City of', 'Enterprise Services (DES), Dept. of'.
    Research queries and narratives need the natural order: 'City of Springfield'."""
    m = INVERTED_ORG.match((name or "").strip())
    return f"{m.group('head')} {m.group('name')}".strip() if m else (name or "").strip()


def load_vendors_downloading(opp_dir):
    """Company names from research/portal-vendors-downloading.csv (bidkit.portal) — the competitive field
    as the portal shows it. [] if absent."""
    import csv
    p = os.path.join(opp_dir, "research", "portal-vendors-downloading.csv")
    if not os.path.exists(p):
        return []
    try:
        with open(p, newline="", encoding="utf-8") as fh:
            return [{"company": r.get("company", ""), "status": r.get("status", "")} for r in csv.DictReader(fh) if r.get("company")]
    except Exception:
        return []


def portal_record_facts(recs):
    """Turn portal records into the same fact shapes the text extractors produce, so they merge into the
    schedule / money / contacts with a clear source. The portal's Date Closed is a real due-date signal
    (no time of day — the RFP text keeps precedence when it states one); the portal's 'Estimated Initial
    Contract Value' is the single best value signal when the agency filled it in."""
    events, money, contacts = [], [], []
    for r in recs:
        src = f"portal-record.json ({r.get('portal_key') or r.get('portal') or 'portal'})"
        for kind, key, label in (("issued", "posted", "Date Posted"), ("due", "closed", "Date Closed")):
            d = r.get(key)
            if d and re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(d)):
                events.append({"kind": kind, "date": d, "time": None, "tz": None, "source": src,
                               "snippet": f"{label} {d} on the portal's solicitation page ({r.get('reference') or r.get('system_id') or ''})".strip()})
        v = r.get("estimated_value")
        if isinstance(v, (int, float)) and v > 0:
            money.append({"kind": "value", "amount": float(v), "text": r.get("estimated_value_text") or f"{v:,.2f}", "source": src,
                          "snippet": r.get("estimated_value_label") or
                          "Estimated Initial Contract Value stated on the portal's solicitation page (the agency's own figure at posting)"})
        c = r.get("contact") or {}
        if c.get("email") or c.get("phone") or c.get("name"):
            contacts.append({"source": src, "emails": [c["email"]] if c.get("email") else [], "phones": [c["phone"]] if c.get("phone") else [],
                             "lines": [f"Contact: {c.get('name') or ''} {c.get('email') or ''} {c.get('phone') or ''}".strip()]})
    return events, money, contacts


# ----------------------------------------------------------------------------- extractors
def _snip(text, start, end, pad=70):
    s = max(0, start - pad); e = min(len(text), end + pad)
    return re.sub(r"\s+", " ", text[s:e]).strip()


def _parse_date(m, kind):
    try:
        if kind == "ymd": y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        elif kind == "mdy":
            mo, d, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if y < 100: y += 2000
        else:
            mo = MONTHS[m.group(1)[:3].lower()]; d = int(m.group(2)); y = int(m.group(3))
        return dt.date(y, mo, d)
    except (ValueError, KeyError):
        return None


SENT_END = re.compile(r"[.;!?]\s|\n\s*\n|\n(?=[A-Z])")


def sentence_bounds(text, start, end, max_back=220, max_fwd=160):
    """(s, e) of the sentence containing [start, end): bounded by ./;/blank line/newline-before-capital.
    Classifying a date by ITS OWN sentence (not a fixed window) stops the previous sentence's cue
    ("Pre-proposal conference …") from being attributed to the next date."""
    s = max(0, start - max_back)
    seg = text[s:start]
    ends = [m.end() for m in SENT_END.finditer(seg)]
    s = s + ends[-1] if ends else s
    e_lim = min(len(text), end + max_fwd)
    m = SENT_END.search(text, end, e_lim)
    e = m.start() + 1 if m else e_lim
    return s, e


def find_dated_events(text, source):
    """Every date in the text, classified by the cue word in ITS OWN sentence (cue before the date wins,
    then cue after it), with an optional time + timezone from the same sentence. DETECTED candidates."""
    events = []
    for pat, kind in DATE_PATTERNS:
        for m in pat.finditer(text):
            d = _parse_date(m, kind)
            if not d or d.year < 2000 or d.year > 2100:
                continue
            s, e = sentence_bounds(text, m.start(), m.end())
            before = text[s:m.start()]; after = text[m.end():e]
            ekind = "other"
            for name, cue in EVENT_CUES:
                if cue.search(before):
                    ekind = name; break
            if ekind == "other":
                for name, cue in EVENT_CUES:
                    if cue.search(after):
                        ekind = name; break
            t = TIME_RE.search(after) or TIME_RE.search(before); tz = TZ_RE.search(after) or TZ_RE.search(before)
            time_s = None
            if t:
                hh = int(t.group(1)) % 12 + (12 if t.group(3).lower() == "p" else 0)
                time_s = f"{hh:02d}:{t.group(2) or '00'}"
            events.append({"kind": ekind, "date": d.isoformat(), "time": time_s,
                           "tz": tz.group(1) if tz else None, "snippet": _snip(text, m.start(), m.end()),
                           "source": source})
    return events


def find_money(text, source):
    out = []
    for m in MONEY_RE.finditer(text):
        raw = m.group(1).replace(",", ""); mult = (m.group(3) or "").lower()
        try:
            val = float(raw + ("." + m.group(2) if m.group(2) else ""))
        except ValueError:
            continue
        val *= {"million": 1e6, "mm": 1e6, "m": 1e6, "k": 1e3, "thousand": 1e3, "billion": 1e9}.get(mult, 1)
        s, e = sentence_bounds(text, m.start(), m.end())
        ctx = text[s:e]                      # the figure's OWN sentence, so a neighbouring NTE clause can't relabel it
        kind = "other"
        for name, cue in MONEY_CUES:
            if cue.search(ctx):
                kind = name; break
        out.append({"kind": kind, "amount": val, "text": m.group(0).strip(), "snippet": _snip(text, m.start(), m.end()), "source": source})
    return out


def find_contacts(text, source):
    emails = sorted(set(EMAIL_RE.findall(text)))
    phones = sorted(set(re.sub(r"\s+", " ", p) for p in PHONE_RE.findall(text)))
    lines = []
    for ln in text.splitlines():
        s = ln.strip()
        if 6 < len(s) < 200 and CONTACT_LINE_RE.search(s) and (EMAIL_RE.search(s) or PHONE_RE.search(s) or re.search(r"[A-Z][a-z]+ [A-Z][a-z]+", s)):
            lines.append(s)
    return {"emails": emails, "phones": phones, "lines": lines[:12], "source": source}


def find_snippets(text, source, cue_re, need_re=None, window=160, limit=20):
    """The SENTENCE containing each cue_re hit (and need_re if given), deduped."""
    out, seen = [], set()
    for m in cue_re.finditer(text):
        s, e = sentence_bounds(text, m.start(), m.end(), max_back=window, max_fwd=window)
        snip = re.sub(r"\s+", " ", text[s:e]).strip()
        if need_re and not need_re.search(snip):
            continue
        key = snip[:80].lower()
        if key in seen or len(snip) < 12:
            continue
        seen.add(key); out.append({"snippet": snip, "source": source})
        if len(out) >= limit: break
    return out


def find_evaluation(text, source, limit=12):
    """Sentences with points/percent that sit under an evaluation heading: the cue may be in the same
    sentence OR within ~400 chars before it (headings like 'EVALUATION CRITERIA' precede their lists)."""
    out, seen = [], set()
    for m in EVAL_RE.finditer(text):
        s, e = sentence_bounds(text, m.start(), m.end())
        if not (EVAL_CUE_RE.search(text[s:e]) or EVAL_CUE_RE.search(text[max(0, s - 400):s])):
            continue
        snip = re.sub(r"\s+", " ", text[s:e]).strip()
        key = snip[:80].lower()
        if key in seen or len(snip) < 8: continue
        seen.add(key); out.append({"snippet": snip, "source": source})
        if len(out) >= limit: break
    return out


def find_attachments(text, source):
    """Referenced attachment labels with the most common following title."""
    refs = {}
    for m in ATTACH_RE.finditer(text):
        kind = m.group(1).title()
        tok = re.sub(r"^(No\.?\s*)", "", m.group(2)).upper()
        label = f"{kind} {tok}"
        e = refs.setdefault(label, {"label": label, "count": 0, "titles": {}, "sources": set()})
        e["count"] += 1; e["sources"].add(source)
        if m.group(3):
            t = m.group(3).strip(" -–—:")
            e["titles"][t] = e["titles"].get(t, 0) + 1
    out = []
    for e in refs.values():
        title = max(e["titles"], key=e["titles"].get) if e["titles"] else None
        out.append({"label": e["label"], "count": e["count"], "title": title, "sources": sorted(e["sources"])})
    return out


def match_attachments_to_files(refs, rfp_dir):
    """Mark each referenced attachment present/missing by fuzzy filename match on the label token
    ("Attachment C2" ~ "C2 - ...", "Exhibit B" ~ "*Exhibit B*") and on the title words."""
    files = []
    if os.path.isdir(rfp_dir):
        for root, _, fs in os.walk(rfp_dir):
            files += [f for f in fs if f.lower().endswith(DOC_EXT) and not f.startswith("portal-page-")]
    low = [f.lower() for f in files]
    for r in refs:
        kind, _, tok = r["label"].partition(" ")
        tok_l = tok.lower()
        hits = [files[i] for i, f in enumerate(low)
                if re.search(rf"(^|[^a-z0-9]){re.escape(tok_l)}([^a-z0-9]|$)", f) or f"{kind.lower()} {tok_l}" in f]
        if not hits and r.get("title"):
            words = [w for w in re.findall(r"[a-z0-9]{4,}", r["title"].lower()) if w not in TITLE_STOP][:3]
            if words:
                hits = [files[i] for i, f in enumerate(low) if all(w in f for w in words)]
        r["files"] = sorted(set(hits)); r["present"] = bool(hits)
    return refs


def detect_format(snippets):
    """Best-effort structured format rules from the collected format snippets (all DETECTED)."""
    joined = " ".join(s["snippet"] for s in snippets)
    fmt = {}
    m = re.search(r"\b(Arial|Times New Roman|Calibri|Helvetica|Verdana|Garamond|Cambria)\b", joined, re.I)
    if m: fmt["font_family"] = m.group(1).title() if m.group(1).lower() != "times new roman" else "Times New Roman"
    m = re.search(r"(?:no (?:smaller|less) than|minimum(?: of)?|at least)\s*[“\"]?\s*(\d{1,2})\s*[-–]?\s*(?:point|pt)", joined, re.I) \
        or re.search(r"(\d{1,2})\s*[-–]?\s*(?:point|pt)\b", joined, re.I)
    if m: fmt["min_font_pt"] = int(m.group(1))
    m = re.search(r"margins?\D{0,40}?(\d(?:\.\d+)?|1/2|½|3/4|¾)\s*[-–]?\s*(?:inch|in\b|\"|”)", joined, re.I)
    if m:
        v = {"1/2": 0.5, "½": 0.5, "3/4": 0.75, "¾": 0.75}.get(m.group(1), None)
        fmt["margin_in"] = v if v is not None else float(m.group(1))
    # page limit: several statements may exist (cover letter ≤2 pp, body ≤24 numbered pp). Prefer the one
    # that says "numbered pages", else take the LARGEST stated cap; accept "twenty-four (24)" word forms.
    cands = []
    for mm in re.finditer(r"(?:maximum of|not (?:to )?exceed|limited to|no more than|up to)\s*([A-Za-z-]+\s*)?\(?(\d{1,3})\)?\s*(numbered\s*)?pages", joined, re.I):
        cands.append((int(mm.group(2)), bool(mm.group(3))))
    for mm in re.finditer(r"(?:maximum of|not (?:to )?exceed|limited to|no more than|up to)\s*([a-z-]+)\s*(numbered\s*)?pages", joined, re.I):
        n = NUMWORDS.get(mm.group(1).lower())
        if n: cands.append((n, bool(mm.group(2))))
    if cands:
        numbered = [n for n, nb in cands if nb]
        fmt["page_limit"] = max(numbered) if numbered else max(n for n, _ in cands)
    if re.search(r"8\s?[½1/2]\s?[\"”x×]|8\.5\s?[\"”x×]|letter", joined, re.I): fmt["page_size"] = "letter"
    m = re.search(r"(\d{1,3})\s*mb\b", joined, re.I)
    if m: fmt["file_size_cap_mb"] = int(m.group(1))
    return fmt


# ----------------------------------------------------------------------------- research plan
SYSTEM_PHRASES = re.compile(r"\b((?:records|jail|corrections|case|content|document|permit|asset|fleet|learning|student|financial|enterprise resource|utility billing|court|property|evidence|dispatch|computer[- ]aided dispatch|human resources|payroll|tax|land|work order|customer information|library|recreation|inmate|offender|parking|911|CAD|RMS|JMS|CMS|ERP|SIS|LMS|CIS|GIS|HRIS)\s+(?:management\s+)?(?:information\s+)?(?:system|software|solution|platform)s?)\b", re.I)


def system_keywords_from_text(text, limit=3):
    """System-type phrases ('records management system', 'jail management system') by frequency in the
    document text — the terms that actually find prior awards and incumbents in web search."""
    counts = {}
    for m in SYSTEM_PHRASES.finditer(text):
        k = re.sub(r"\s+", " ", m.group(1).lower()).replace("softwares", "software")
        counts[k] = counts.get(k, 0) + 1
    return [k for k, _ in sorted(counts.items(), key=lambda kv: -kv[1])[:limit]]


AGENCY_RE = re.compile(r"\b((?:City|Town|Village|County|Borough|Port|Housing Authority|Public Utility District|Transit Authority)\s+of\s+[A-Z][A-Za-z.'-]+(?:\s+[A-Z][A-Za-z.'-]+){0,3}"
                       r"|[A-Z][A-Za-z.'-]+(?:\s+[A-Z][A-Za-z.'-]+){0,3}\s+(?:County|Parish|School District|Public Schools|Fire District|Water District|Utility District|Transit|Housing Authority|College|Community College|University|State University|Police Department|Sheriff'?s? Office|Regional Library|Port District|Health District|PUD)"
                       r"|(?:State|Commonwealth|Department|Office|Board)\s+of\s+[A-Z][A-Za-z.'&-]+(?:\s+[A-Z][A-Za-z.'&-]+){0,4})\b")
GENERIC_AGENCY = re.compile(r"^(procurement|purchasing|contracts?|finance|business)\b.*(office|department|division|services)?$|^(office|department|division)\b", re.I)
AGENCY_STOP_LEAD = re.compile(r"^(The|This|Any|Each|Such|Said|Request|Notice|Attn|Attention|Dear)\s", re.I)


def detect_agency(text, limit=1):
    """Most frequent agency-shaped phrase in the document text ('City of Example', 'Example Valley
    College', 'Example County'). Returns a string or None. Frequency beats first-mention: the buyer's
    own name appears on every page; vendors' and partners' names appear once."""
    counts = {}
    for m in AGENCY_RE.finditer(text):
        a = re.sub(r"\s+", " ", m.group(1)).strip(" .,;:")
        if AGENCY_STOP_LEAD.match(a) or len(a) < 6: continue
        counts[a] = counts.get(a, 0) + 1
    if not counts: return None
    return sorted(counts.items(), key=lambda kv: (-kv[1], len(kv[0])))[0][0]


def clean_agency(agency):
    """'City of Example / CONSORTIUM' -> 'City of Example'; drop trailing parentheticals; a generic
    department name ('Procurement Office') is not an agency and returns ''."""
    a = (agency or "").split("/")[0].split("—")[0].split(" - ")[0]
    a = re.sub(r"\s*\(.*?\)\s*", " ", a).strip()
    return "" if GENERIC_AGENCY.match(a) else a


def research_queries(agency, title, number, keywords):
    """Concrete web-search queries + what each is for. Anchored on the lesson that the incumbent's
    prior contract (found in a council packet) validated the estimate on a prior bid."""
    agency = clean_agency(agency)
    a = f'"{agency}"' if agency else ""
    n = f'"{number}"' if number else ""
    sys_kw = " ".join(f'"{k}"' for k in keywords[:2]) if keywords else (f'"{title}"' if title else "")
    Q = []
    def q(query, why, look_for):
        query = re.sub(r"\s+", " ", query).strip()
        if query: Q.append({"query": query, "why": why, "look_for": look_for})
    q(f"{a} {n}", "the solicitation itself elsewhere (aggregators, prior postings, cancelled/reissued versions)",
      "prior versions, Q&A, planholder lists, bid tabulation from a previous round")
    q(f"{a} {sys_kw} contract award", "who has the CURRENT contract and what it cost",
      "award amount, term, vendor name (the incumbent), award date")
    q(f"{a} {sys_kw} council agenda OR commission agenda OR board agenda", "council/commission packets approving the incumbent's contract or this procurement",
      "staff report with budget, contract amount, term, vendor; often a PDF with the full agreement attached")
    q(f"{a} {sys_kw} staff report OR agenda bill OR resolution", "the approving resolution for the last contract/renewal",
      "not-to-exceed amount, annual maintenance, renewal history")
    q(f"site:legistar.com {a} {sys_kw}", "Legistar-hosted agendas (many WA/US cities)", "legislation text + attachments with contract $")
    q(f"{a} {sys_kw} (civicclerk OR granicus OR civicweb OR municode) ", "other agenda-hosting platforms", "same as above")
    q(f"{a} {sys_kw} RFP {dt.date.today().year - 8}..{dt.date.today().year - 1}", "the PREVIOUS solicitation for this system",
      "the incumbent's winning proposal, bid tabulation, evaluation scores, pricing pages")
    q(f"{a} \"notice of award\" OR \"intent to award\" {sys_kw}", "award notices", "winning vendor + amount")
    q(f"{a} budget {sys_kw} \"capital\" OR \"CIP\" OR \"technology\"", "adopted budget / CIP line for this project", "appropriated amount and fiscal year")
    q(f"{a} \"public records\" request contract {sys_kw}", "how to request the incumbent contract if not online", "records-request portal/email")
    q(f"site:des.wa.gov {n} OR ({a} {sys_kw})", "Washington DES/WEBS postings and award history", "award notice, planholders")
    q(f"{sys_kw} RFP award {dt.date.today().year - 3}..{dt.date.today().year} \"per year\" OR \"total cost\"", "COMPARABLE awards at other agencies (price benchmarks)",
      "contract values for similar scope/size; vendor GSA or state price lists")
    q(f"{sys_kw} \"GSA schedule\" OR \"price list\" OR \"state contract\" pricing", "published vendor pricing", "unit prices, annual subscription/maintenance rates")
    return Q


def keywords_from_title(title):
    if not title: return []
    t = re.sub(r"\b(re|fw|fwd)\s*:\s*", "", title, flags=re.I)
    t = re.sub(r"#?\b[A-Z]{0,4}-?\d{2,4}-\d{2,4}(?:-[A-Z0-9]+)?\b", " ", t)
    words = [w for w in re.findall(r"[A-Za-z][A-Za-z0-9/&-]+", t) if w.lower() not in TITLE_STOP and len(w) > 2]
    # keep phrases: consecutive capitalised words
    phrases, cur = [], []
    for w in words:
        if w[0].isupper() or w.isupper():
            cur.append(w)
        else:
            if len(cur) >= 2: phrases.append(" ".join(cur))
            cur = []
    if len(cur) >= 2: phrases.append(" ".join(cur))
    return phrases[:3] or words[:3]


RESEARCH_PLAN_HEADER = """# Research plan — {oid}

_Generated by bidkit.brief on {when}. Run these searches during intake (Kiro can run them with its web
search tool: "run the research plan for {oid}"). Record what you find in `findings.md` with the source
URL for every number. The goal is to VALIDATE the estimated value and know the incumbent before pricing._

## Why this matters
On a prior bid the incumbent's previous contract (found in a city council packet via web search) gave the
actual prior award amount and annual maintenance — the single best check on our estimate. Every
solicitation gets this treatment.

## What to establish
1. **Incumbent** — who runs the current system, since when, and on what contract.
2. **Prior award value and term** — from the approving council/board resolution or contract (PDF).
3. **Prior solicitation** — the last RFP for this system, its bid tabulation / evaluation scores, and the
   winning proposal if public.
4. **Budget signals** — adopted budget / CIP line, grant funding, NTE language in this RFP
   (see opportunity-facts.md → Money).
5. **Comparable awards** — same system type at similar-size agencies in the last 3 years; vendor price lists.
6. **Agency context** — recent news (consolidations, audits, failed implementations) that explains the ask.

## Queries
"""

FINDINGS_TEMPLATE = """# Research findings — {oid}

_Fill during intake. Every figure needs a source URL. Mark unknowns UNKNOWN, never guess._

## Incumbent
- Vendor / product: UNKNOWN
- Since (year): UNKNOWN
- Source:

## Prior contract
- Award amount (one-time): UNKNOWN
- Annual maintenance / subscription: UNKNOWN
- Term / renewals: UNKNOWN
- Approving body + date: UNKNOWN
- Source (council packet / contract PDF):

## Prior solicitation
- RFP number / year: UNKNOWN
- Bid tabulation / scores: UNKNOWN
- Winning proposal available? UNKNOWN
- Source:

## Budget signals
- From this RFP (see opportunity-facts.md → Money): 
- Adopted budget / CIP / grant: UNKNOWN
- Source:

## Comparable awards (other agencies, last ~3 years)
| Agency | System / vendor | Award $ | Annual $ | Term | Year | Source |
|---|---|---|---|---|---|---|
|  |  |  |  |  |  |  |

## Estimated value — working view
- Low / likely / high: UNKNOWN (basis: )
- What would change it:

## Agency context / notes
"""


# ----------------------------------------------------------------------------- assembly
def build_facts(opp_dir, max_snippets=12):
    sources = gather_sources(opp_dir)
    facts = {"opportunity": os.path.basename(os.path.abspath(opp_dir)), "generated": dt.datetime.now(dt.timezone.utc).isoformat(),
             "documents": [], "events": [], "money": [], "contacts": [], "format_snippets": [], "attachments": [],
             "evaluation": [], "submission": [], "setaside": [], "term": [], "insurance": []}
    all_text = []
    for label, path in sources:
        text, meta = extract_text(path)
        meta["label"] = label; facts["documents"].append(meta)
        if not text.strip(): continue
        all_text.append(text)
        evs = find_dated_events(text, label)
        if label.lower().endswith((".xlsx", ".xlsm", ".xls")):
            # spreadsheet cells are data, not deadline sentences: a Q&A log's row timestamps and a
            # reference workbook's example go-live date must never become schedule items. Keep them
            # visible as "other" mentions only.
            for e in evs: e["kind"] = "other"
        facts["events"] += evs
        facts["money"] += find_money(text, label)
        c = find_contacts(text, label)
        if c["emails"] or c["phones"] or c["lines"]: facts["contacts"].append(c)
        facts["format_snippets"] += find_snippets(text, label, FORMAT_RE, limit=max_snippets)
        facts["attachments"] += find_attachments(text, label)
        facts["evaluation"] += find_evaluation(text, label, limit=max_snippets)
        facts["submission"] += find_snippets(text, label, SUBMIT_RE, need_re=SUBMIT_HOW_RE, limit=max_snippets)
        facts["setaside"] += find_snippets(text, label, SETASIDE_RE, limit=6)
        facts["term"] += find_snippets(text, label, TERM_RE, limit=6)
        facts["insurance"] += find_snippets(text, label, re.compile(r"\binsurance\b", re.I), need_re=MONEY_RE, limit=6)
    # the portal's own record (bidkit.portal fetch): dates, estimated value, contact — same shapes, own source
    facts["portal_records"] = load_portal_records(opp_dir)
    pev, pmoney, pcontacts = portal_record_facts(facts["portal_records"])
    facts["events"] += pev; facts["money"] += pmoney; facts["contacts"] += pcontacts
    # merge attachments across sources
    merged = {}
    for r in facts["attachments"]:
        e = merged.setdefault(r["label"], {"label": r["label"], "count": 0, "title": None, "sources": []})
        e["count"] += r["count"]; e["sources"] = sorted(set(e["sources"] + r["sources"]))
        if r["title"] and not e["title"]: e["title"] = r["title"]
    facts["attachments"] = match_attachments_to_files(sorted(merged.values(), key=lambda x: x["label"]), os.path.join(opp_dir, "rfp"))
    # dedupe events on (kind, date), keeping the RICHEST mention (with time, then tz) — the email
    # usually names the date bare while the RFP gives "2:00 PM Pacific Time"
    best = {}
    for e in facts["events"]:
        k = (e["kind"], e["date"])
        score = (e["time"] is not None, e["tz"] is not None, len(e["snippet"]))
        if k not in best or score > best[k][0]:
            best[k] = (score, e)
    facts["events"] = sorted((v[1] for v in best.values()), key=lambda x: (x["date"], x["kind"]))
    facts["format_detected"] = detect_format(facts["format_snippets"])
    facts["agency_detected"] = detect_agency("\n".join(all_text))
    # the portal names the issuing organization explicitly — better than a regex over the text
    portal_org = next((normalize_org_name(r["organization"]) for r in facts["portal_records"] if r.get("organization")), None)
    if portal_org:
        facts["agency_detected"] = portal_org
    facts["vendors_downloading"] = load_vendors_downloading(opp_dir)
    # Schedule picks, per kind. Prose sources (pdf/docx/email) outrank spreadsheet cells: an xlsx
    # Q&A log stringifies every row's datetime as "2026-07-24 00:00:00", which is data, not a
    # deadline sentence — on a real RFP that made every kind land on the Q&A date. Within a source
    # class prefer a mention WITH a time (real deadlines have one), then the earliest date.
    def rank(e):
        prose = not e["source"].lower().endswith((".xlsx", ".xlsm", ".xls"))
        return (0 if prose else 1, 0 if e["time"] else 1, e["date"])
    picks = {}
    for e in sorted(facts["events"], key=rank):
        if e["kind"] != "other" and e["kind"] not in picks:
            picks[e["kind"]] = e
    facts["schedule"] = picks
    facts["money_value_signals"] = sorted([m for m in facts["money"] if m["kind"] == "value"], key=lambda m: -m["amount"])[:10]
    return facts, "\n".join(all_text)


def facts_markdown(facts, oj):
    F = facts; L = []
    L += [f"# Opportunity facts — {F['opportunity']}", "",
          f"_Generated {F['generated']} by bidkit.brief from {len(F['documents'])} document(s). Every item below is **DETECTED** by "
          f"pattern-matching the solicitation text and shown with its source snippet — nothing is confirmed until a human reads the RFP. "
          f"The summary and estimated value are produced in the intake step (`run intake on …`) using this sheet plus `research/findings.md`._", ""]
    ag = oj.get("agency") or "(unknown)"
    ag_src = oj.get("agency_source") or ""
    ag_note = f" — _{ag_src}_" if ag_src else ""
    if F.get("agency_detected") and F["agency_detected"] != oj.get("agency"):
        ag_note += f"; detected in documents: **{F['agency_detected']}**"
    L += ["## Identity", f"- **Title:** {oj.get('title') or '(unknown)'}", f"- **Agency:** {ag}{ag_note}",
          f"- **Source:** {(oj.get('source') or {}).get('type', '?')} — {(oj.get('source') or {}).get('from', '')}", ""]
    if F.get("portal_records"):
        L += ["## Portal record (DETECTED from the portal's own solicitation page — the agency's stated figures)"]
        for r in F["portal_records"]:
            c = r.get("contact") or {}
            L += [f"- **{r.get('title') or '(untitled)'}** — {normalize_org_name(r.get('organization'))} · `{r.get('portal_key') or r.get('portal')}` [page]({r.get('url')})",
                  f"  - reference # **{r.get('reference') or '—'}** (portal id {r.get('system_id') or '—'}); posted {r.get('posted') or '—'}; **closes {r.get('closed') or '—'}** (time of day not shown on the portal — take it from the RFP)",
                  f"  - estimated initial contract value: **{r.get('estimated_value_text') or 'blank on the portal'}**",
                  f"  - contact: {c.get('name') or '—'} {c.get('email') or ''} {c.get('phone') or ''}".rstrip(),
                  f"  - {len(r.get('documents') or [])} document(s), {len(r.get('amendments') or [])} amendment(s): "
                  + "; ".join([d['name'] for d in (r.get('documents') or [])] + [f"{a.get('date') or ''} {a['name']}".strip() for a in (r.get('amendments') or [])])]
        if F.get("vendors_downloading"):
            vd = F["vendors_downloading"]
            L += [f"- **Vendors downloading (portal, opt-in list): {len(vd)}** — " + ", ".join(v["company"] for v in vd[:25]) + (" …" if len(vd) > 25 else "")
                  + " → `research/portal-vendors-downloading.csv`"]
        L.append("")
    L += ["## Documents read", "| Document | Type | Pages/Sheets | Words |", "|---|---|---|---|"]
    for d in F["documents"]:
        L.append(f"| {d['label']} | {d['type']} | {d.get('pages', d.get('sheets', ''))} | {d.get('words', 0):,}{' — ERROR ' + d['error'] if d.get('error') else ''} |")
    L.append("")
    L += ["## Schedule (DETECTED — earliest date per kind; confirm every one in the RFP)", "| Kind | Date | Time | TZ | Evidence |", "|---|---|---|---|---|"]
    order = ["issued", "prebid", "questions", "addenda", "due", "interviews", "award", "start"]
    for k in order:
        e = F["schedule"].get(k)
        if e: L.append(f"| **{k}** | {e['date']} | {e['time'] or ''} | {e['tz'] or ''} | _{e['snippet'][:140]}_ ({e['source']}) |")
    others = [e for e in F["events"] if e["kind"] == "other"]
    if others:
        L += ["", f"<details><summary>{len(others)} other dated mention(s)</summary>", ""]
        L += [f"- {e['date']} — _{e['snippet'][:120]}_ ({e['source']})" for e in others[:25]] + ["", "</details>"]
    L.append("")
    L += ["## Money (DETECTED)", ""]
    if F["money_value_signals"]:
        L += ["**Value signals** (budget / NTE / estimate / award cues):", "| Amount | Text | Evidence |", "|---|---|---|"]
        L += [f"| ${m['amount']:,.0f} | {m['text']} | _{m['snippet'][:150]}_ ({m['source']}) |" for m in F["money_value_signals"]]
        L.append("")
    other_money = [m for m in F["money"] if m["kind"] != "value"]
    if other_money:
        L += [f"<details><summary>{len(other_money)} other dollar mention(s) (bond / insurance / fees / other)</summary>", ""]
        L += [f"- [{m['kind']}] {m['text']} — _{m['snippet'][:120]}_ ({m['source']})" for m in other_money[:30]] + ["", "</details>", ""]
    if not F["money"]:
        L += ["_No dollar figures found in the documents read. Value must come from research (research/)._", ""]
    L += ["## Contacts (DETECTED)"]
    for c in F["contacts"]:
        L.append(f"- **{c['source']}** — emails: {', '.join(c['emails']) or '—'}; phones: {', '.join(c['phones']) or '—'}")
        L += [f"  - _{ln[:160]}_" for ln in c["lines"][:6]]
    L.append("")
    L += ["## Format rules (DETECTED — scored; confirm and copy into opportunity.json → format)"]
    if F["format_detected"]:
        L.append("- Parsed: " + ", ".join(f"{k}={v}" for k, v in F["format_detected"].items()))
    L += [f"- _{s['snippet'][:180]}_ ({s['source']})" for s in F["format_snippets"][:12]] or ["- none found"]
    L.append("")
    L += ["## Attachments / forms referenced (DETECTED — map by TITLE against rfp/; ✔ present, ✘ missing)", "| Label | Title (as referenced) | Mentions | In rfp/ |", "|---|---|---|---|"]
    for r in F["attachments"]:
        mark = "✔ " + ", ".join(r["files"][:2]) if r.get("present") else "✘ **missing**"
        L.append(f"| {r['label']} | {r['title'] or ''} | {r['count']} | {mark} |")
    if not F["attachments"]: L.append("| — | none referenced | | |")
    L.append("")
    L += ["## Submission (DETECTED)"] + [f"- _{s['snippet'][:200]}_ ({s['source']})" for s in F["submission"][:10]] + [""]
    L += ["## Evaluation criteria (DETECTED — lines with points/percent near 'evaluation')"] + [f"- _{s['snippet'][:200]}_ ({s['source']})" for s in F["evaluation"][:12]] + [""]
    if F["term"]: L += ["## Term (DETECTED)"] + [f"- _{s['snippet'][:180]}_ ({s['source']})" for s in F["term"][:6]] + [""]
    if F["insurance"]: L += ["## Insurance (DETECTED)"] + [f"- _{s['snippet'][:180]}_ ({s['source']})" for s in F["insurance"][:6]] + [""]
    if F["setaside"]: L += ["## Set-aside / preference language (DETECTED)"] + [f"- _{s['snippet'][:180]}_ ({s['source']})" for s in F["setaside"][:6]] + [""]
    L += ["## Next", "1. Confirm the schedule and format rules against the RFP text; set `due` (and `format`) in `opportunity.json` with `confirmed: true`.",
          "2. Run `research/research-plan.md` and fill `research/findings.md` (incumbent, prior award $, comparables).",
          "3. In Kiro: **run intake on this opportunity** — the summary and estimated value are written there, citing this sheet and the findings.", ""]
    return "\n".join(L)


def update_opportunity_json(opp_dir, facts):
    p = os.path.join(opp_dir, "opportunity.json")
    if not os.path.exists(p): return False
    with open(p, encoding="utf-8") as fh:
        oj = json.load(fh)
    changed = False
    sched = facts["schedule"]
    DETECTED = "bidkit.brief (DETECTED)"
    # Ownership rule: a value is refreshable if a bidkit tool wrote it (source starts with "bidkit.")
    # and it is not confirmed — that covers our own earlier run AND the inbox poller's cruder guess from
    # the email text. Human-entered (no source) or confirmed values are never touched. Without this, a
    # bad first detection (e.g. a spreadsheet cell date, or the email's earliest date) would stick forever.
    def ours(entry):
        return isinstance(entry, dict) and str(entry.get("source", "")).startswith("bidkit.") and not entry.get("confirmed")
    # agency: the inbox poller guesses it from the sender's display name ("Procurement Office"); prefer a
    # real agency name found in the documents, and record the source so the human can see it's a guess.
    if str(oj.get("agency_source", "")).startswith("bidkit.") or not oj.get("agency"):
        better = facts.get("agency_detected")
        if better and better != oj.get("agency"):
            oj["agency"] = better; oj["agency_source"] = DETECTED; changed = True
    due = oj.setdefault("due", {})
    if sched.get("due") and not due.get("confirmed") and (not due.get("date") or ours(due)):
        new = {"date": sched["due"]["date"], "time": sched["due"].get("time") or "", "confirmed": False, "source": DETECTED}
        if sched["due"].get("tz"): new["tz_hint"] = sched["due"]["tz"]
        if due.get("timezone"): new["timezone"] = due["timezone"]
        if {k: due.get(k) for k in new} != new:
            due.update(new); changed = True
    elif not sched.get("due") and ours(due):
        oj["due"] = {k: v for k, v in due.items() if k in ("timezone",)}; oj["due"].update({"date": "", "time": "", "confirmed": False}); changed = True
    ms = oj.setdefault("milestones", {})
    for k in ("questions", "prebid", "addenda", "award", "start", "interviews", "issued"):
        cur = ms.get(k)
        if sched.get(k):
            new = {"date": sched[k]["date"], "time": sched[k].get("time"), "confirmed": False, "source": DETECTED}
            if not cur or (ours(cur) and {kk: cur.get(kk) for kk in new} != new):
                ms[k] = new; changed = True
        elif ours(cur):
            del ms[k]; changed = True                       # our stale detection, no longer supported by the text
    if facts.get("format_detected") and oj.get("format_detected") != facts["format_detected"]:
        oj["format_detected"] = facts["format_detected"]; changed = True
    recs = facts.get("portal_records") or []
    if recs:
        r = recs[-1]
        summary = {"portal": r.get("portal_key") or r.get("portal"), "url": r.get("url"), "system_id": r.get("system_id"),
                   "reference": r.get("reference"), "title": r.get("title"), "organization": r.get("organization"),
                   "posted": r.get("posted"), "closed": r.get("closed"), "estimated_value": r.get("estimated_value"),
                   "estimated_value_text": r.get("estimated_value_text"), "documents": len(r.get("documents") or []),
                   "amendments": len(r.get("amendments") or []), "source": DETECTED}
        if oj.get("portal_record") != summary:
            oj["portal_record"] = summary; changed = True
        # the agency's own number, if the poller couldn't read one off the subject
        num = str(oj.get("solicitation_number") or "").strip()
        if (not num or re.fullmatch(r"\{\{.*\}\}", num)) and r.get("reference"):
            oj["solicitation_number"] = r["reference"]; oj["solicitation_number_source"] = DETECTED; changed = True
    oj["facts"] = {"file": "opportunity-facts.md", "generated": facts["generated"], "documents": len(facts["documents"])}
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(oj, fh, indent=2, ensure_ascii=False)
    return changed


# a due date is only AUTO-CONFIRMABLE from a document source carrying an explicit submission-deadline
# cue — not a mere "closing/closes" mention, which also appears in boilerplate. These are the phrases
# that unambiguously mark the proposal submission deadline in an RFP.
_DEADLINE_CUE = re.compile(
    r"\b(proposals?\s+(?:due|deadline|must be received)|submission\s+deadline|"
    r"deadline\s+for\s+(?:submi\w+|proposals?|receipt)|bids?\s+due|due\s+date|"
    r"no later than|must be received (?:by|no later)|closing\s+(?:date\s+and\s+)?time)\b", re.I)


def _is_document_source(source):
    """True if a schedule pick's source is a solicitation DOCUMENT in rfp/ — not the email, a follow-up
    email, the portal record, or a spreadsheet cell (which brief already demotes to 'other')."""
    s = (source or "").lower()
    if s in ("email",) or s.startswith("followup:") or "portal-record" in s:
        return False
    return s.endswith((".pdf", ".docx", ".doc", ".txt", ".md", ".html", ".htm"))


def auto_confirm_due_from_docs(opp_dir, facts):
    """DETERMINISTIC due-date confirmation, provenance-gated. Promote due.confirmed=true ONLY when the
    chosen due event was detected from a DOCUMENT in rfp/ (not the email/portal), has a TIME, and its
    snippet carries an explicit submission-deadline cue. Email-only guesses, timeless mentions, and
    bare 'closing' boilerplate stay unconfirmed for a human. Returns (confirmed_bool, note)."""
    p = os.path.join(opp_dir, "opportunity.json")
    if not os.path.exists(p):
        return False, "no opportunity.json"
    with open(p, encoding="utf-8") as fh:
        oj = json.load(fh)
    due = oj.get("due") or {}
    if due.get("confirmed"):
        return False, "already confirmed"
    pick = (facts.get("schedule") or {}).get("due")
    if not pick:
        return False, "no due detection"
    src, snippet, time_s = pick.get("source"), pick.get("snippet", ""), pick.get("time")
    if not _is_document_source(src):
        return False, f"due source is not a document ({src})"
    if not time_s:
        return False, "no time of day detected (ambiguous — leave for human)"
    if not _DEADLINE_CUE.search(snippet or ""):
        return False, "no explicit submission-deadline cue in the snippet"
    # all gates passed — confirm with a citation naming the document and the quoted deadline sentence
    due["date"] = pick["date"]
    due["time"] = time_s
    if pick.get("tz"):
        due["tz_hint"] = pick["tz"]
    due["confirmed"] = True
    due["source"] = f"bidkit.brief auto-confirm from {src}: \"{(snippet or '').strip()[:140]}\""
    oj["due"] = due
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(oj, fh, indent=2, ensure_ascii=False)
    return True, f"confirmed {due['date']} {time_s} from {src}"


def write_research(opp_dir, oj, facts, all_text=""):
    rd = os.path.join(opp_dir, "research"); os.makedirs(rd, exist_ok=True)
    oid = os.path.basename(os.path.abspath(opp_dir))
    recs = facts.get("portal_records") or []
    prec = recs[-1] if recs else {}
    number = str(oj.get("solicitation_number") or "").strip() or prec.get("reference") or \
        ((oid.split("-")[0] + "-" + oid.split("-")[1]) if re.match(r"^[A-Z]{0,4}-?\d{2,4}-\d{2,4}", oid) else None)
    if number and re.fullmatch(r"\{\{.*\}\}", number): number = None
    # the portal's title is the solicitation's real name; the email subject (opportunity.json title) may
    # be a notification wrapper around it
    title = prec.get("title") or oj.get("title") or oid.replace("-", " ")   # older opportunity.json files have no title
    # system-type phrases from the DOCUMENTS beat title words: "records management system" finds the
    # incumbent's prior award; a folder slug / acronym alone finds nothing
    kws = system_keywords_from_text(all_text) or keywords_from_title(title)
    facts["research_keywords"] = kws
    # agency for queries: a human-set agency wins; else the agency detected in the documents; else the
    # poller's sender-name guess (which clean_agency() drops if it is just a department name)
    agency = oj.get("agency") if not str(oj.get("agency_source", "")).startswith("bidkit.") else (facts.get("agency_detected") or oj.get("agency") or "")
    facts["research_agency"] = clean_agency(agency)
    qs = research_queries(agency, title, number, kws)
    lines = [RESEARCH_PLAN_HEADER.format(oid=oid, when=facts["generated"])]
    for i, q in enumerate(qs, 1):
        lines += [f"{i}. `{q['query']}`", f"   - why: {q['why']}", f"   - look for: {q['look_for']}", ""]
    lines += ["## Portal / aggregator checks", "- The agency's own procurement page (planholders, addenda, Q&A).",
              "- Public Purchase / BidNet / WEBS record for this number (award history, prior postings).",
              "- WEBS advanced search (SearchBid.aspx) with status <ALL> and the agency name: the agency's CLOSED/ARCHIVED "
              "solicitations show what it bought before and often name the incumbent; documents there may still download.",
              "- State contract databases for the system type (piggyback-able contracts set price ceilings).", ""]
    if prec:
        lines += ["## Portal record (already fetched — cite it)",
                  f"- {prec.get('title')} — {normalize_org_name(prec.get('organization'))}; ref # {prec.get('reference')}; posted {prec.get('posted')}, closes {prec.get('closed')}",
                  f"- Estimated Initial Contract Value on the portal: **{prec.get('estimated_value_text') or 'blank'}** — "
                  + ("the agency's own figure; anchor the value estimate on it and explain any departure." if prec.get("estimated_value") else
                     "the agency left it blank; the value must come from research."), ""]
    vd = facts.get("vendors_downloading") or []
    if vd:
        lines += [f"## Competitive field — {len(vd)} vendor(s) the portal lists as having downloaded this solicitation",
                  "_Opt-in list only (a floor). Search each likely competitor with the system type for prior awards and pricing; "
                  "look for the incumbent among them. Contact details: research/portal-vendors-downloading.csv._", ""]
        lines += [f"- {v['company']}" + (f" ({v['status']})" if v.get("status") else "") for v in vd[:60]] + [""]
    lines += ["## Money signals already in the RFP"]
    lines += [f"- ${m['amount']:,.0f} — _{m['snippet'][:120]}_ ({m['source']})" for m in facts["money_value_signals"][:8]] or ["- none detected"]
    with open(os.path.join(rd, "research-plan.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    fp = os.path.join(rd, "findings.md")
    if not os.path.exists(fp):
        with open(fp, "w", encoding="utf-8") as fh:
            fh.write(FINDINGS_TEMPLATE.format(oid=oid))
    with open(os.path.join(rd, "queries.json"), "w", encoding="utf-8") as fh:
        json.dump(qs, fh, indent=2)
    return len(qs)


def run(opp_dir, research=True, max_snippets=12):
    opp_dir = os.path.abspath(opp_dir)
    ojp = os.path.join(opp_dir, "opportunity.json")
    oj = json.load(open(ojp, encoding="utf-8")) if os.path.exists(ojp) else {}
    facts, all_text = build_facts(opp_dir, max_snippets=max_snippets)
    nq = write_research(opp_dir, oj, facts, all_text) if research else 0
    with open(os.path.join(opp_dir, "opportunity-facts.md"), "w", encoding="utf-8") as fh:
        fh.write(facts_markdown(facts, oj))
    with open(os.path.join(opp_dir, "opportunity-facts.json"), "w", encoding="utf-8") as fh:
        json.dump(facts, fh, indent=2, ensure_ascii=False, default=str)
    changed = update_opportunity_json(opp_dir, facts)
    return facts, changed, nq


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("opportunity"); ap.add_argument("--json", action="store_true")
    ap.add_argument("--no-research", action="store_true"); ap.add_argument("--max-snippets", type=int, default=12)
    ap.add_argument("--confirm-from-docs", action="store_true",
                    help="auto-confirm the due date when it was detected from a document in rfp/ with a time and an explicit deadline cue (provenance-gated)")
    a = ap.parse_args(argv)
    if not os.path.isdir(a.opportunity):
        print(f"not a directory: {a.opportunity}"); return 2
    facts, changed, nq = run(a.opportunity, research=not a.no_research, max_snippets=a.max_snippets)
    confirmed, confirm_note = (False, "")
    if a.confirm_from_docs:
        confirmed, confirm_note = auto_confirm_due_from_docs(a.opportunity, facts)
    if a.json:
        print(json.dumps({"schedule": facts["schedule"], "value_signals": facts["money_value_signals"],
                          "format_detected": facts["format_detected"], "attachments": facts["attachments"],
                          "documents": facts["documents"]}, indent=2, default=str)); return 0
    docs = facts["documents"]
    print(f"brief: {len(docs)} document(s), {sum(d.get('words', 0) for d in docs):,} words -> opportunity-facts.md")
    for k in ("issued", "prebid", "questions", "due", "award"):
        e = facts["schedule"].get(k)
        if e: print(f"   {k:10s} {e['date']} {e['time'] or ''} {e['tz'] or ''}  (DETECTED)")
    if facts["money_value_signals"]:
        print("   value signals: " + ", ".join(f"${m['amount']:,.0f}" for m in facts["money_value_signals"][:5]))
    missing = [r["label"] for r in facts["attachments"] if not r.get("present")]
    if missing: print(f"   attachments referenced but NOT in rfp/: {', '.join(missing[:8])}")
    if facts["format_detected"]: print(f"   format detected: {facts['format_detected']}")
    print(f"   opportunity.json {'updated' if changed else 'unchanged'}; research plan: {nq} queries -> research/")
    if a.confirm_from_docs:
        print(f"   due auto-confirm: {'CONFIRMED — ' + confirm_note if confirmed else 'not confirmed (' + confirm_note + ')'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
