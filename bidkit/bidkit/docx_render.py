#!/usr/bin/env python3
"""
docx_render.py — render Markdown to .docx (python-docx), with a colored-ANSWER convention for
partner intake checklists.

Why: bid partners (bidwriters) want a completed checklist .docx with our answers in a distinct
color (green by default; some ask for red), and a proposal narrative as .docx when they
assemble the final document. Hand-editing .docx files is not regenerable — after every fact
change the doc drifts. Rendering from Markdown keeps the source diffable and re-runnable.

Markdown supported: #/##/### headings, - and * bullets, 1. numbered lists, | tables | (with
:--- / ---: alignment), **bold**, *italic*, --- rules, > blockquotes, paragraphs, and the
checklist conventions below. HTML comments are stripped.

Checklist conventions:
  >> [TAG] answer text        an ANSWER line, rendered in --answer-color; TAG (ANSWERED, CONFIRM,
                              DECISION NEEDED, BUILD SCOPE, OPEN, or any) rendered bold
  >> answer text              answer without a tag
  Everything else is the partner's original text, rendered in black.

CLI:
  python3 -m bidkit.docx_render <in.md> <out.docx> [--answer-color green|red|blue|#RRGGBB]
        [--font Calibri] [--size 11] [--cfg opportunity.json]
  (--cfg is optional: if given and format.font_family_allow exists, its first family is the
   default font, and body_font_pt the default size)
"""
import argparse, os, re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAMED_COLORS = {"green": (0x00, 0x80, 0x00), "red": (0xC0, 0x00, 0x00), "blue": (0x00, 0x33, 0xCC),
                "black": (0, 0, 0), "grey": (0x55, 0x55, 0x55), "gray": (0x55, 0x55, 0x55)}
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "center": WD_ALIGN_PARAGRAPH.CENTER}


def parse_color(spec):
    spec = (spec or "green").strip().lower()
    if spec in NAMED_COLORS:
        return RGBColor(*NAMED_COLORS[spec])
    m = re.fullmatch(r"#?([0-9a-f]{6})", spec)
    if not m:
        raise ValueError(f"bad color {spec!r}")
    h = m.group(1)
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def add_runs(paragraph, text, color=None, bold=False, italic=False):
    """Render inline **bold** / *italic* into runs; optional forced color/bold/italic."""
    i, cur_b, cur_i, buf = 0, bold, italic, ""
    def flush():
        nonlocal buf
        if buf:
            r = paragraph.add_run(buf); r.bold = cur_b; r.italic = cur_i
            if color is not None: r.font.color.rgb = color
        buf = ""
    while i < len(text):
        if text.startswith("**", i):
            flush(); cur_b = not cur_b if not bold else True; i += 2; continue
        if text[i] == "*" and (i + 1 < len(text) and text[i + 1] != " "):
            flush(); cur_i = not cur_i if not italic else True; i += 1; continue
        buf += text[i]; i += 1
    flush()


def parse_align(sep_cells):
    out = []
    for c in sep_cells:
        c = c.strip()
        out.append("center" if c.startswith(":") and c.endswith(":") else "right" if c.endswith(":") else "left")
    return out


def cells_of(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def render(md_text, out_path, answer_color="green", font="Calibri", size=11):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = font; st.font.size = Pt(size)
    ans_rgb = parse_color(answer_color)
    md_text = re.sub(r"<!--.*?-->", "", md_text, flags=re.DOTALL)
    lines = md_text.splitlines(); i = 0; n_answers = 0
    while i < len(lines):
        raw = lines[i]; s = raw.strip()
        if not s:
            i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            doc.add_heading(re.sub(r"\*", "", m.group(2)), level=len(m.group(1))); i += 1; continue
        if s == "---":
            p = doc.add_paragraph(); p.add_run("_" * 60).font.color.rgb = RGBColor(0x99, 0x99, 0x99); i += 1; continue
        # ANSWER line
        m = re.match(r"^>>\s*(?:\[([^\]]+)\]\s*)?(.*)", s)
        if m:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(18)
            if m.group(1):
                r = p.add_run(f"[{m.group(1).strip()}] "); r.bold = True; r.font.color.rgb = ans_rgb
            add_runs(p, m.group(2), color=ans_rgb); n_answers += 1; i += 1; continue
        if s.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote" if "Intense Quote" in [x.name for x in doc.styles] else None)
            add_runs(p, s[2:], italic=True); i += 1; continue
        if s.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            header = cells_of(s); aligns = parse_align(cells_of(lines[i + 1])); rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(cells_of(lines[i])); i += 1
            ncol = max(len(header), *(len(r) for r in rows)) if rows else len(header)
            t = doc.add_table(rows=1 + len(rows), cols=ncol); t.style = "Table Grid"
            # NOTE: do not set cell.text = "" first — python-docx leaves a phantom empty run.
            for ci, c in enumerate(header + [""] * (ncol - len(header))):
                add_runs(t.cell(0, ci).paragraphs[0], c, bold=True)
            for ri, row in enumerate(rows, start=1):
                for ci, c in enumerate(row + [""] * (ncol - len(row))):
                    cell = t.cell(ri, ci)
                    add_runs(cell.paragraphs[0], c)
                    if ci < len(aligns): cell.paragraphs[0].alignment = ALIGN[aligns[ci]]
            doc.add_paragraph(); continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Number"); add_runs(p, m.group(2)); i += 1; continue
        if s.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, s[2:]); i += 1; continue
        # paragraph with soft-wrapped continuation lines
        text = s; i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||- |\* |\d+\.\s|---$|>)", lines[i].strip()):
            text += " " + lines[i].strip(); i += 1
        add_runs(doc.add_paragraph(), text)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc.save(out_path)
    return {"answers": n_answers, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("md"); ap.add_argument("out")
    ap.add_argument("--answer-color", default="green")
    ap.add_argument("--font"); ap.add_argument("--size", type=float)
    ap.add_argument("--cfg")
    a = ap.parse_args(argv)
    font, size = a.font or "Calibri", a.size or 11
    if a.cfg:
        from . import config
        fmt = config.load(a.cfg)["format"]
        font = a.font or (fmt.get("font_family_allow") or ["Calibri"])[0]
        size = a.size or fmt.get("body_font_pt", 11)
    with open(a.md, encoding="utf-8") as fh:
        stats = render(fh.read(), a.out, answer_color=a.answer_color, font=font, size=size)
    print(f"rendered {a.out}: {stats}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
