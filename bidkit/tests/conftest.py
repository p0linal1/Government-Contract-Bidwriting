"""
Shared pytest fixtures for bidkit. All fixtures are built programmatically so the test suite
carries no opportunity-specific data. Fixtures mimic the structural features of real agency
workbooks/forms: uniform requirement rows, section-header rows with blank numbers, an
UNNUMBERED real requirement, merged cells, x14 dropdown validations, and flat PDF forms with
underscore signature lines.
"""
import json, os, sys, zipfile, re
import pytest

# make `import bidkit` work when running pytest from the bidkit/ directory or repo root
HERE = os.path.dirname(os.path.abspath(__file__))
PKG_ROOT = os.path.dirname(HERE)
if PKG_ROOT not in sys.path:
    sys.path.insert(0, PKG_ROOT)

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
import fitz


# --------------------------------------------------------------------------- xlsx fixtures
def _build_requirements_xlsx(path):
    """A requirement workbook: Control tab (vocab), Instructions tab, two requirement tabs.
    Rows: header at 2; section headers with blank #; one UNNUMBERED real requirement; and a
    merged-cell row to exercise the merge guard. We then inject an x14 dataValidation by
    editing the sheet XML, because openpyxl cannot write x14 extensions natively."""
    wb = openpyxl.Workbook()
    ctl = wb.active; ctl.title = "Control"
    vocab = ["Meets Requirement", "Exceeds Requirement", "Requires Modification - Cost",
             "Requires Modification - No Cost", "Do not meet requirement"]
    ctl["A1"] = "Allowed Responses"
    for i, v in enumerate(vocab, start=2):
        ctl[f"A{i}"] = v

    ins = wb.create_sheet("Instructions"); ins["A1"] = "Fill column C from the dropdown."

    for tab in ("General", "Reporting"):
        ws = wb.create_sheet(tab)
        ws["A1"] = tab
        ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "#", "Requirement", "Response", "Comments"
        ws["B3"] = "General Requirements"            # section header (blank #)
        ws["A4"], ws["B4"] = 1, "Ability to log in with MFA."
        ws["A5"], ws["B5"] = 2, "Provide an audit trail of all edits."
        # UNNUMBERED but real requirement (author forgot the number)
        ws["B6"] = "Provide an inmate contact report showing all other inmates assigned to the same cell or pod as a specified inmate over a date range."
        ws["A7"], ws["B7"] = 3, "Ability to export reports to PDF."
        ws["B8"] = "Searches and Queries"            # another header
        ws["A9"], ws["B9"] = 4, "Provide a quick search across all inmates."
        # a merged-cell row (form-style) — C10:D10 merged; only C10 is writable
        ws["A10"], ws["B10"] = 5, "Describe your escalation process."
        ws.merge_cells("C10:D10")
    wb.save(path)

    # inject an x14 dataValidation extLst into each requirement sheet's XML
    with zipfile.ZipFile(path) as z:
        names = z.namelist(); contents = {n: z.read(n) for n in names}
    x14 = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
           'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
           '<x14:dataValidations xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" count="1">'
           '<x14:dataValidation type="list" allowBlank="1" showInputMessage="1" showErrorMessage="1">'
           '<x14:formula1><xm:f>Control!$A$2:$A$6</xm:f></x14:formula1>'
           '<xm:sqref>C3:C200</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
    for n in names:
        if re.match(r"xl/worksheets/sheet\d+\.xml", n):
            xml = contents[n].decode("utf-8")
            # only requirement sheets have a "Requirement" header cell; cheap check via sharedStrings is
            # awkward, so inject into every sheet except the first two (Control, Instructions).
            idx = int(re.search(r"sheet(\d+)\.xml", n).group(1))
            if idx >= 3:
                xml = xml.replace("</worksheet>", x14 + "</worksheet>")
                contents[n] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names: z.writestr(n, contents[n])
    return path


@pytest.fixture
def req_xlsx(tmp_path):
    return _build_requirements_xlsx(str(tmp_path / "req.xlsx"))


@pytest.fixture
def opp_dir(tmp_path, req_xlsx):
    """An opportunity folder with opportunity.json pointing at the fixture workbook."""
    opp = tmp_path / "OPP-001"; opp.mkdir()
    (opp / "rfp").mkdir()
    os.replace(req_xlsx, str(opp / "rfp" / "req.xlsx"))
    cfg = {
        "id": "OPP-001",
        "format": {"font_family_allow": ["Arial"], "min_font_pt": 10, "body_font_pt": 11,
                   "margin_in": 0.75, "page_size": "letter", "page_limit": 24,
                   "numbered_from_heading": "PART I",
                   "exclude_numbering_headings": ["Cover Letter", "Resumes"],
                   "header_footer_band_pt": 40},
        "fonts": {"regular": "/System/Library/Fonts/Supplemental/Arial.ttf",
                  "bold": "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
                  "italic": "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
                  "bold_italic": "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"},
        "workbooks": {"R": {"path": "rfp/req.xlsx", "header_row": 2,
                            "cols": {"num": 1, "req": 2, "resp": 3, "comment": 4},
                            "skip_tabs": ["Instructions", "Control"],
                            "control_tab": "Control", "control_range": "A2:A6",
                            "placeholders": ["Y / N / NA", "If Yes include"]}},
        "forms": {"F": {"path": "rfp/form.pdf", "page": 1,
                        "fields": {"printed_name": "Printed Name", "company": "Company Name",
                                   "signature": "Signature", "date": "Date"},
                        "signature_note": "(original signature on submission)"}},
        "signer": {"printed": "Test Person, Principal", "company": "Test LLC",
                   "typed_signature": "Test Person", "signature_image": "", "date": ""},
        "redact": {"domain": "example.com", "keep_contact_for": ["Test Person"]},
        "folders": {"extract": "deliverables/bid/extract", "draft": "deliverables/bid/workbooks-draft",
                    "final": "deliverables/bid/final", "submit": "../OPP-001-SUBMIT"},
    }
    (opp / "opportunity.json").write_text(json.dumps(cfg, indent=2))
    _build_flat_form_pdf(str(opp / "rfp" / "form.pdf"))
    return str(opp)


# --------------------------------------------------------------------------- pdf fixtures
def _build_flat_form_pdf(path):
    """A flat certification form: labels above underscore lines, like agency forms."""
    d = fitz.open(); p = d.new_page(width=612, height=792)
    p.insert_text((72, 100), "CERTIFICATION AND ASSURANCES", fontsize=14, fontname="helv")
    p.insert_text((72, 140), "I certify the statements above are true.", fontsize=11, fontname="helv")
    # name / company line
    p.insert_text((72, 220), "_" * 34, fontsize=11, fontname="helv")
    p.insert_text((360, 220), "_" * 30, fontsize=11, fontname="helv")
    p.insert_text((72, 234), "Printed Name & Title", fontsize=9, fontname="helv")
    p.insert_text((360, 234), "Company Name", fontsize=9, fontname="helv")
    # signature / date line
    p.insert_text((72, 300), "_" * 34, fontsize=11, fontname="helv")
    p.insert_text((360, 300), "_" * 30, fontsize=11, fontname="helv")
    p.insert_text((72, 314), "Signature", fontsize=9, fontname="helv")
    p.insert_text((360, 314), "Date", fontsize=9, fontname="helv")
    d.save(path); d.close()
    return path


@pytest.fixture
def flat_form_pdf(tmp_path):
    return _build_flat_form_pdf(str(tmp_path / "form.pdf"))


@pytest.fixture
def blank_image_pdf(tmp_path):
    """A PDF whose page is only an image with no text — the 'Preview flattened it' failure."""
    path = str(tmp_path / "blank.pdf")
    d = fitz.open(); p = d.new_page(width=612, height=792)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 50, 50), False); pix.clear_with(255)
    p.insert_image(fitz.Rect(0, 0, 612, 792), pixmap=pix)
    d.save(path); d.close()
    return path


@pytest.fixture
def resume_pdf(tmp_path):
    """A resume with personal contact info to scrub."""
    path = str(tmp_path / "resume.pdf")
    d = fitz.open(); p = d.new_page(width=612, height=792)
    p.insert_text((72, 80), "Jane Doe", fontsize=16, fontname="helv")
    p.insert_text((72, 100), "Springfield, USA | jane.doe.personal@gmail.com | (555) 123-4567", fontsize=10, fontname="helv")
    p.insert_text((72, 116), "linkedin.com/in/janedoe | github.com/janedoe | https://janedoe.dev", fontsize=10, fontname="helv")
    p.insert_text((72, 150), "Professional Summary", fontsize=12, fontname="helv")
    p.insert_text((72, 170), "Senior engineer with 8 years of experience.", fontsize=10, fontname="helv")
    d.save(path); d.close()
    return path


@pytest.fixture
def sample_md(tmp_path):
    md = """# Cover Letter

Dear Evaluators,

We are pleased to submit this proposal.

<!-- PAGEBREAK -->

# PART I

## 1. General Information

- Item one with **bold** and *italic*.
- Item two.

| Col A | Col B |
|---|---|
| a | b |

<!-- PAGEBREAK -->

# Resumes

## Jane Doe

Senior engineer.

<!-- PAGEBREAK -->

# PART IV

Pricing follows.
"""
    p = tmp_path / "doc.md"; p.write_text(md); return str(p)


# --------------------------------------------------------------------------- real-sample lookup
REPO_ROOT = os.path.dirname(PKG_ROOT)


def real_sample(*relpath):
    """Path to a real solicitation file used for regression tests, or None if absent.
    Looks under the repo root (legacy layout) and under opportunities/ (current layout), so the
    regression tests keep running after an opportunity is archived."""
    rel = os.path.join(*relpath)
    for base in (REPO_ROOT, os.path.join(REPO_ROOT, "opportunities")):
        p = os.path.join(base, rel)
        if os.path.exists(p):
            return p
    return None
