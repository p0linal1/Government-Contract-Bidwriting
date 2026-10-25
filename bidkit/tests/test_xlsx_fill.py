import json, os, warnings
warnings.filterwarnings("ignore")
import openpyxl
from bidkit import config, xlsx_extract, xlsx_fill
from bidkit._xlsx_common import read_package, sheet_xml_map, count_validations, out_of_order_rows


def _opp(opp_dir):
    return config.load(os.path.join(opp_dir, "opportunity.json"))


def _x14_total(path):
    names, contents = read_package(path)
    return sum(count_validations(contents[p].decode())[0] for p in sheet_xml_map(contents).values())


def test_fill_preserves_x14_and_writes_values(opp_dir):
    opp = _opp(opp_dir)
    rows, _ = xlsx_extract.extract_workbook(opp, "R")
    for r in rows:
        r["response"] = "Meets Requirement"
        r["comment"] = f"Comment for {r['tab']} #{r['num']}"
    res, dst = xlsx_fill.fill_responses(opp, "R", rows)
    assert res.ok, res.report()
    # 10 responses + 10 comments, minus the 2 comments whose D cell is merged into C (row 10
    # on each tab) — those are SKIPPED (informational), not refused, because the response
    # anchor holds the row's value.
    assert res.written == 18
    assert sorted(res.skipped_comments) == [("General", "D10"), ("Reporting", "D10")]
    assert res.refused_merged == []
    src = opp.resolve(opp.workbook("R")["path"])
    assert _x14_total(src) == 2 and _x14_total(dst) == 2       # one per requirement tab, preserved
    assert res.validations_before == res.validations_after == 2

    wb = openpyxl.load_workbook(dst, data_only=True)
    ws = wb["Reporting"]
    assert ws["C4"].value == "Meets Requirement"
    assert ws["D4"].value == "Comment for Reporting #1"
    assert ws["C6"].value is None                              # unnumbered row untouched by responses mode


def test_fill_never_produces_out_of_order_rows(opp_dir):
    opp = _opp(opp_dir)
    rows, _ = xlsx_extract.extract_workbook(opp, "R")
    for r in rows:
        r["response"] = "Exceeds Requirement"; r["comment"] = "x"
    res, dst = xlsx_fill.fill_responses(opp, "R", rows)
    names, contents = read_package(dst)
    for p in sheet_xml_map(contents).values():
        assert out_of_order_rows(contents[p].decode()) == []
    assert res.out_of_order_after == 0


def test_bad_vocabulary_is_refused(opp_dir):
    opp = _opp(opp_dir)
    rows, _ = xlsx_extract.extract_workbook(opp, "R")
    rows[0]["response"] = "Totally Meets It"                  # not in Control!A2:A6
    rows[1]["response"] = "Meets Requirement"
    res, dst = xlsx_fill.fill_responses(opp, "R", rows)
    assert not res.ok
    assert len(res.bad_vocab) == 1 and res.bad_vocab[0][2] == "Totally Meets It"
    wb = openpyxl.load_workbook(dst, data_only=True)
    assert wb[rows[0]["tab"]].cell(rows[0]["row"], 3).value is None   # refused, not written
    assert wb[rows[1]["tab"]].cell(rows[1]["row"], 3).value == "Meets Requirement"


def test_merged_away_write_is_refused_and_anchor_allowed(opp_dir):
    opp = _opp(opp_dir)
    src = opp.resolve(opp.workbook("R")["path"])
    dst = os.path.join(opp_dir, "out.xlsx")
    # C10:D10 is merged in the fixture. D10 is merged-away; C10 is the anchor.
    res = xlsx_fill.fill_cells(src, dst, {"General": {"D10": "hidden!", "C10": "visible answer"}})
    assert res.written == 1
    assert len(res.refused_merged) == 1 and res.refused_merged[0][1] == "D10"
    assert not res.ok
    wb = openpyxl.load_workbook(dst, data_only=True)
    assert wb["General"]["C10"].value == "visible answer"
    assert wb["General"]["D10"].value is None


def test_cells_mode_cli(opp_dir, tmp_path):
    cells = {"General": {"C4": "Meets Requirement", "D4": "via cells mode"}}
    p = tmp_path / "cells.json"; p.write_text(json.dumps(cells))
    rc = xlsx_fill.main([os.path.join(opp_dir, "opportunity.json"), "R", "--cells", str(p)])
    assert rc == 0
    dst = os.path.join(opp_dir, "deliverables", "bid", "workbooks-draft", "req.xlsx")
    wb = openpyxl.load_workbook(dst, data_only=True)
    assert wb["General"]["D4"].value == "via cells mode"


def test_original_is_never_modified(opp_dir):
    opp = _opp(opp_dir)
    src = opp.resolve(opp.workbook("R")["path"])
    before = open(src, "rb").read()
    rows, _ = xlsx_extract.extract_workbook(opp, "R")
    for r in rows: r["response"] = "Meets Requirement"
    xlsx_fill.fill_responses(opp, "R", rows)
    assert open(src, "rb").read() == before
