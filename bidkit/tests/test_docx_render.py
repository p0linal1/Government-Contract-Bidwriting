import os
from docx import Document
from docx.shared import RGBColor
from bidkit import docx_render

CHECKLIST = """# Intake Checklist

## 1. Company

Is the vendor registered in SAM.gov?
>> [ANSWERED] Yes — active registration.

Provide UEI.
>> [CONFIRM] On file; reconfirm currency before submission.

Which components will you bid?
>> [DECISION NEEDED] Full scope vs. RMS-only — options attached.

Free-form note.
>> No tag on this one.

## 2. Table

| Item | Status |
|---|---:|
| A | **done** |
| B | *open* |

- bullet one
- bullet two

1. first
2. second

> a quoted line

---
"""


def test_answers_are_colored_and_tags_bold(tmp_path):
    out = str(tmp_path / "c.docx")
    stats = docx_render.render(CHECKLIST, out, answer_color="green")
    assert stats["answers"] == 4 and stats["tables"] == 1
    d = Document(out)
    green = RGBColor(0x00, 0x80, 0x00)
    answer_paras = [p for p in d.paragraphs if any(r.font.color.rgb == green for r in p.runs if r.font.color and r.font.color.rgb)]
    assert len(answer_paras) == 4
    tagged = [p for p in answer_paras if p.runs and p.runs[0].text.startswith("[")]
    assert len(tagged) == 3 and all(p.runs[0].bold for p in tagged)
    assert any("[DECISION NEEDED]" in p.runs[0].text for p in tagged)
    # partner text is NOT colored
    q = next(p for p in d.paragraphs if p.text.startswith("Is the vendor"))
    assert all(not (r.font.color and r.font.color.rgb) for r in q.runs)


def test_red_answer_color_and_hex(tmp_path):
    out = str(tmp_path / "r.docx")
    docx_render.render(">> [OPEN] pending", out, answer_color="red")
    d = Document(out)
    assert d.paragraphs[0].runs[0].font.color.rgb == RGBColor(0xC0, 0x00, 0x00)
    docx_render.render(">> x", out, answer_color="#112233")
    assert Document(out).paragraphs[0].runs[0].font.color.rgb == RGBColor(0x11, 0x22, 0x33)


def test_inline_bold_italic_and_table_alignment(tmp_path):
    out = str(tmp_path / "t.docx")
    docx_render.render(CHECKLIST, out)
    d = Document(out)
    t = d.tables[0]
    assert t.cell(0, 0).text == "Item" and t.cell(1, 1).text == "done"
    assert t.cell(1, 1).paragraphs[0].runs[0].bold is True
    assert t.cell(2, 1).paragraphs[0].runs[0].italic is True
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    assert t.cell(1, 1).paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    styles = [p.style.name for p in d.paragraphs]
    assert "List Bullet" in styles and "List Number" in styles
    assert any(p.style.name.startswith("Heading") for p in d.paragraphs)


def test_font_and_size_applied(tmp_path):
    out = str(tmp_path / "f.docx")
    docx_render.render("hello", out, font="Arial", size=12)
    d = Document(out)
    assert d.styles["Normal"].font.name == "Arial" and d.styles["Normal"].font.size.pt == 12


def test_cli_with_cfg_uses_allowed_font(opp_dir, tmp_path, capsys):
    md = tmp_path / "in.md"; md.write_text("# T\n\nbody\n>> [ANSWERED] ok\n")
    out = str(tmp_path / "o.docx")
    rc = docx_render.main([str(md), out, "--cfg", os.path.join(opp_dir, "opportunity.json")])
    assert rc == 0 and "answers': 1" in capsys.readouterr().out
    assert Document(out).styles["Normal"].font.name == "Arial"
